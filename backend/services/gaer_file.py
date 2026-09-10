import io
import math
import re
import unicodedata
from collections import OrderedDict, deque

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


MAX_GAER_FILE_SIZE = 10 * 1024 * 1024
EAN_HEADERS = {"CODICEABARRE", "EAN", "EAN13", "BARCODE"}
QUANTITY_HEADERS = {"QUANTITA", "QTA", "QTY", "QUANTITY"}
ARTICLE_HEADERS = {"ARTICOLO", "CODICEARTICOLO", "ITEM", "ITEMCODE"}


class GaerFileError(ValueError):
    pass


class GaerExportError(ValueError):
    pass


def _normalized_header(value) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9]", "", normalized.upper())


def _header_row(sheet):
    candidates = []
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25), values_only=True),
        start=1,
    ):
        labels = [str(value or "").strip() for value in row]
        normalized = [_normalized_header(value) for value in labels]
        exact_matches = sum(
            value in EAN_HEADERS or value in QUANTITY_HEADERS
            for value in normalized
        )
        populated = sum(bool(value) for value in labels)
        candidates.append((exact_matches, populated, -row_number, row_number, labels))
    if not candidates:
        raise GaerFileError("Il foglio Excel non contiene righe leggibili.")
    return max(candidates)[3:]


def _column_options(labels):
    return [
        {
            "key": get_column_letter(index),
            "label": label or f"Colonna {get_column_letter(index)}",
            "index": index,
        }
        for index, label in enumerate(labels, start=1)
        if label
    ]


def _find_column(options, aliases):
    return next(
        (
            option["key"]
            for option in options
            if _normalized_header(option["label"]) in aliases
        ),
        None,
    )


def _ean(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        result = str(value)
    elif isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return ""
        result = str(int(value))
    else:
        result = str(value).strip().replace(" ", "")
        if result.endswith(".0") and result[:-2].isdigit():
            result = result[:-2]
    return result if result.isdigit() and 8 <= len(result) <= 13 else ""


def _quantity(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        parsed = float(str(value).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None
    if not math.isfinite(parsed) or parsed <= 0:
        return None
    return parsed


def parse_gaer_workbook(
    content: bytes,
    *,
    ean_column: str | None = None,
    quantity_column: str | None = None,
    header_row: int | None = None,
) -> dict:
    if not content:
        raise GaerFileError("Il file Gaer è vuoto.")
    if len(content) > MAX_GAER_FILE_SIZE:
        raise GaerFileError("Il file Gaer supera il limite di 10 MB.")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise GaerFileError("Il file non è un Excel .xlsx valido.") from error

    try:
        sheet = next((item for item in workbook.worksheets if item.max_row > 0), None)
        if sheet is None:
            raise GaerFileError("Il file Gaer non contiene fogli leggibili.")
        if header_row is not None:
            if header_row < 1 or header_row > min(sheet.max_row, 25):
                raise GaerFileError("La riga delle intestazioni selezionata non è valida.")
            labels = [
                str(cell.value or "").strip()
                for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row))
            ]
        else:
            header_row, labels = _header_row(sheet)
        options = _column_options(labels)
        option_keys = {option["key"] for option in options}
        detected_ean = _find_column(options, EAN_HEADERS)
        detected_quantity = _find_column(options, QUANTITY_HEADERS)
        selected_ean = str(ean_column or detected_ean or "").upper()
        selected_quantity = str(quantity_column or detected_quantity or "").upper()

        if (
            selected_ean not in option_keys
            or selected_quantity not in option_keys
            or selected_ean == selected_quantity
        ):
            return {
                "mapping_required": True,
                "sheet_name": sheet.title,
                "header_row": header_row,
                "columns": options,
                "suggested_mapping": {
                    "ean": detected_ean,
                    "quantity": detected_quantity,
                },
            }

        ean_index = next(item["index"] for item in options if item["key"] == selected_ean)
        quantity_index = next(
            item["index"] for item in options if item["key"] == selected_quantity
        )
        description_index = next(
            (
                item["index"]
                for item in options
                if _normalized_header(item["label"]) in {"DESCRIZIONE", "DESCRIZIONEARTICOLO"}
            ),
            None,
        )
        article_index = next(
            (
                item["index"]
                for item in options
                if _normalized_header(item["label"]) in ARTICLE_HEADERS
            ),
            None,
        )
        stock = OrderedDict()
        warnings = []
        parsed_rows = 0
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            ean_value = row[ean_index - 1] if len(row) >= ean_index else None
            quantity_value = row[quantity_index - 1] if len(row) >= quantity_index else None
            if ean_value in (None, "") and quantity_value in (None, ""):
                continue
            normalized_ean = _ean(ean_value)
            quantity = _quantity(quantity_value)
            if not normalized_ean:
                warnings.append({"row": row_number, "message": "EAN non valido o mancante."})
                continue
            if quantity is None:
                warnings.append({
                    "row": row_number,
                    "ean": normalized_ean,
                    "message": "Quantità non valida o non positiva.",
                })
                continue
            description = ""
            if description_index and len(row) >= description_index:
                description = str(row[description_index - 1] or "").strip()
            article = ""
            if article_index and len(row) >= article_index:
                article = str(row[article_index - 1] or "").strip()
            item = stock.setdefault(normalized_ean, {
                "ean": normalized_ean,
                "article": article,
                "description": description,
                "quantity": 0.0,
                "rows": [],
            })
            item["quantity"] += quantity
            item["rows"].append(row_number)
            if not item["description"] and description:
                item["description"] = description
            if not item["article"] and article:
                item["article"] = article
            parsed_rows += 1
        if not stock:
            raise GaerFileError("Nessuna riga Gaer valida trovata nelle colonne selezionate.")
        return {
            "mapping_required": False,
            "sheet_name": sheet.title,
            "header_row": header_row,
            "mapping": {"ean": selected_ean, "quantity": selected_quantity},
            "stock": list(stock.values()),
            "warnings": warnings,
            "parsed_rows": parsed_rows,
        }
    finally:
        workbook.close()


def export_gaer_workbook(
    content: bytes,
    *,
    allocations: dict[str, list[int]],
    ean_column: str,
    quantity_column: str,
    header_row: int,
    sheet_name: str | None = None,
) -> bytes:
    """Write proposed order IDs into free cells without changing source data."""
    if not content:
        raise GaerExportError("Il file Gaer è vuoto.")
    if len(content) > MAX_GAER_FILE_SIZE:
        raise GaerExportError("Il file Gaer supera il limite di 10 MB.")
    try:
        workbook = load_workbook(io.BytesIO(content))
    except Exception as error:
        raise GaerExportError("Il file non è un Excel .xlsx valido.") from error

    try:
        sheet = (
            workbook[sheet_name]
            if sheet_name and sheet_name in workbook.sheetnames
            else next((item for item in workbook.worksheets if item.max_row > 0), None)
        )
        if sheet is None:
            raise GaerExportError("Il file Gaer non contiene fogli leggibili.")
        if header_row < 1 or header_row > sheet.max_row:
            raise GaerExportError("La riga delle intestazioni non è valida.")
        try:
            ean_index = openpyxl_column_index(ean_column)
            quantity_index = openpyxl_column_index(quantity_column)
        except ValueError as error:
            raise GaerExportError("La mappatura delle colonne Gaer non è valida.") from error
        if ean_index == quantity_index:
            raise GaerExportError("EAN e quantità devono usare colonne diverse.")

        last_header_column = max(
            (
                cell.column
                for cell in sheet[header_row]
                if cell.value not in (None, "")
            ),
            default=max(ean_index, quantity_index),
        )
        output_start = last_header_column + 1
        pending = {
            ean: deque(order_ids)
            for ean, order_ids in allocations.items()
            if order_ids
        }

        written = 0
        for row_number in range(header_row + 1, sheet.max_row + 1):
            normalized_ean = _ean(sheet.cell(row_number, ean_index).value)
            queue = pending.get(normalized_ean)
            if not queue:
                continue
            quantity = _quantity(sheet.cell(row_number, quantity_index).value)
            if quantity is None or not float(quantity).is_integer():
                continue
            writable_units = int(quantity)
            target_column = output_start
            while queue and writable_units > 0:
                while sheet.cell(row_number, target_column).value not in (None, ""):
                    target_column += 1
                sheet.cell(row_number, target_column).value = queue.popleft()
                target_column += 1
                writable_units -= 1
                written += 1

        unresolved = sum(len(order_ids) for order_ids in pending.values())
        if unresolved:
            raise GaerExportError(
                f"Impossibile associare {unresolved} unità alle righe del file Gaer. "
                "Aggiorna l’analisi e riprova."
            )
        if written == 0:
            raise GaerExportError("Non ci sono ordini proposti da esportare.")

        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def openpyxl_column_index(column: str) -> int:
    normalized = str(column or "").strip().upper()
    if not re.fullmatch(r"[A-Z]{1,3}", normalized):
        raise ValueError("Invalid column")
    result = 0
    for char in normalized:
        result = result * 26 + ord(char) - 64
    return result
