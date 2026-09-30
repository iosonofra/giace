import io
import math
import re
import unicodedata
from collections import Counter, defaultdict, deque
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.excel_common import normalize_sku_cell
from backend.models import ImportBatch, WarehouseStock
from backend.services.warehouse_import import load_column_mappings


MAX_STOCK_VERIFICATION_FILE_SIZE = 10 * 1024 * 1024
SKU_HEADERS = {
    "SKU", "CODICE", "CODICESKU", "CODICEARTICOLO", "RIFERIMENTO", "REFERENCE",
}
QUANTITY_HEADERS = {
    "QTATOT", "QTATOTALE", "QUANTITA", "QTA", "QTY", "QUANTITY", "STOCK",
}


class StockVerificationError(ValueError):
    pass


def _excel_sheet_title(value, fallback="Verifica giacenze") -> str:
    title = re.sub(r"[\\[\\]:*?/\\\\]", " ", str(value or "")).strip()
    return (title or fallback)[:31]


def _normalized_header(value) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    normalized = "".join(
        char for char in normalized if not unicodedata.combining(char)
    )
    return re.sub(r"[^A-Z0-9]", "", normalized.upper())


def _normalize_sku(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    return normalize_sku_cell(value).strip().upper()


def _quantity(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        parsed = float(str(value).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None
    if not math.isfinite(parsed) or parsed < 0:
        return None
    return parsed


def _header_row(sheet):
    candidates = []
    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row, 25),
            values_only=True,
        ),
        start=1,
    ):
        labels = [str(value or "").strip() for value in row]
        normalized = [_normalized_header(value) for value in labels]
        exact_matches = sum(
            value in SKU_HEADERS or value in QUANTITY_HEADERS
            for value in normalized
        )
        populated = sum(bool(value) for value in labels)
        candidates.append(
            (exact_matches, populated, -row_number, row_number, labels)
        )
    if not candidates:
        raise StockVerificationError(
            "Il file non contiene righe di intestazione leggibili."
        )
    return max(candidates)[3:]


def _column_options(labels):
    return [
        {
            "key": get_column_letter(index),
            "label": label or f"Colonna {get_column_letter(index)}",
            "index": index,
        }
        for index, label in enumerate(labels, start=1)
        if label
    ]


def _find_column(options, aliases):
    return next(
        (
            option["key"]
            for option in options
            if _normalized_header(option["label"]) in aliases
        ),
        None,
    )


def parse_stock_verification_file(
    content: bytes,
    *,
    sku_column: str | None = None,
    quantity_column: str | None = None,
    header_row: int | None = None,
) -> dict:
    if not content:
        raise StockVerificationError("Il file di verifica è vuoto.")
    if len(content) > MAX_STOCK_VERIFICATION_FILE_SIZE:
        raise StockVerificationError(
            "Il file di verifica supera il limite di 10 MB."
        )
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as error:
        raise StockVerificationError(
            "Il file non è un Excel .xlsx valido."
        ) from error

    try:
        sheet = next(
            (item for item in workbook.worksheets if item.max_row > 0),
            None,
        )
        if sheet is None:
            raise StockVerificationError(
                "Il file non contiene fogli leggibili."
            )
        if header_row is not None:
            if header_row < 1 or header_row > min(sheet.max_row, 25):
                raise StockVerificationError(
                    "La riga delle intestazioni selezionata non è valida."
                )
            labels = [
                str(cell.value or "").strip()
                for cell in next(
                    sheet.iter_rows(
                        min_row=header_row,
                        max_row=header_row,
                    )
                )
            ]
        else:
            header_row, labels = _header_row(sheet)

        options = _column_options(labels)
        option_keys = {option["key"] for option in options}
        detected_sku = _find_column(options, SKU_HEADERS)
        detected_quantity = _find_column(options, QUANTITY_HEADERS)
        selected_sku = str(sku_column or detected_sku or "").upper()
        selected_quantity = str(
            quantity_column or detected_quantity or ""
        ).upper()

        if (
            selected_sku not in option_keys
            or selected_quantity not in option_keys
            or selected_sku == selected_quantity
        ):
            return {
                "mapping_required": True,
                "sheet_name": sheet.title,
                "header_row": header_row,
                "columns": options,
                "suggested_mapping": {
                    "sku": detected_sku,
                    "quantity": detected_quantity,
                },
            }

        sku_index = next(
            item["index"] for item in options
            if item["key"] == selected_sku
        )
        quantity_index = next(
            item["index"] for item in options
            if item["key"] == selected_quantity
        )
        rows = []
        warnings = []
        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=header_row + 1,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            sku_value = row[sku_index - 1] if len(row) >= sku_index else None
            quantity_value = (
                row[quantity_index - 1]
                if len(row) >= quantity_index
                else None
            )
            if sku_value in (None, "") and quantity_value in (None, ""):
                continue
            sku = _normalize_sku(sku_value)
            quantity = _quantity(quantity_value)
            if not sku:
                warnings.append({
                    "type": "invalid_row",
                    "row": row_number,
                    "sku": "",
                    "quantity": "" if quantity_value is None else str(quantity_value),
                    "message": "SKU mancante o non valido.",
                })
                continue
            if quantity is None:
                warnings.append({
                    "type": "invalid_quantity",
                    "row": row_number,
                    "sku": sku,
                    "quantity": "" if quantity_value is None else str(quantity_value),
                    "message": "Quantità mancante, negativa o non numerica.",
                })
                continue
            rows.append({
                "sku": sku,
                "quantity": quantity,
                "row": row_number,
            })

        if not rows:
            raise StockVerificationError(
                "Nessuna riga valida trovata nelle colonne selezionate."
            )
        duplicate_counts = Counter(item["sku"] for item in rows)
        return {
            "mapping_required": False,
            "sheet_name": sheet.title,
            "header_row": header_row,
            "mapping": {
                "sku": selected_sku,
                "quantity": selected_quantity,
            },
            "rows": rows,
            "warnings": warnings,
            "parsed_rows": len(rows),
            "duplicate_skus": sum(
                1 for count in duplicate_counts.values() if count > 1
            ),
        }
    finally:
        workbook.close()


def _active_stock(db):
    batch = (
        db.query(ImportBatch)
        .filter(
            ImportBatch.file_type == "warehouse",
            ImportBatch.is_active.is_(True),
        )
        .order_by(ImportBatch.id.desc())
        .first()
    )
    if batch is None:
        raise StockVerificationError(
            "Non è disponibile una giacenza attiva. Sincronizza prima Google Sheets."
        )
    rows = (
        db.query(WarehouseStock)
        .filter(WarehouseStock.import_batch_id == batch.id)
        .order_by(WarehouseStock.id.asc())
        .all()
    )
    if not rows:
        raise StockVerificationError(
            "La giacenza attiva non contiene righe esportabili."
        )
    return batch, rows


def build_stock_verification(db, parsed_file: dict) -> dict:
    batch, stock_rows = _active_stock(db)
    pending = defaultdict(deque)
    for item in parsed_file["rows"]:
        pending[item["sku"]].append(item)

    rows = []
    anomalies = list(parsed_file.get("warnings", []))
    matched_rows = 0
    retained_rows = 0
    reference_rows = 0

    for position, stock_item in enumerate(stock_rows, start=1):
        is_spacer = stock_item.sku.startswith("__spacer_")
        if is_spacer:
            rows.append({
                "position": position,
                "sku": "",
                "description": "",
                "lotto": "",
                "quantity": None,
                "google_quantity": None,
                "source": "spacer",
                "source_row": None,
                "is_spacer": True,
            })
            continue

        reference_rows += 1
        sku = _normalize_sku(stock_item.sku)
        source_item = pending[sku].popleft() if pending[sku] else None
        if source_item:
            quantity = source_item["quantity"]
            source = "file"
            source_row = source_item["row"]
            matched_rows += 1
        else:
            quantity = float(stock_item.qty_total or 0)
            source = "google_sheets"
            source_row = None
            retained_rows += 1
            anomalies.append({
                "type": "missing_in_file",
                "row": None,
                "sku": sku,
                "quantity": quantity,
                "message": (
                    "SKU assente nel file: mantenuta la quantità Google Sheets."
                ),
            })
        rows.append({
            "position": position,
            "sku": stock_item.sku,
            "description": stock_item.description or "",
            "lotto": stock_item.lotto or "",
            "quantity": quantity,
            "google_quantity": float(stock_item.qty_total or 0),
            "source": source,
            "source_row": source_row,
            "is_spacer": False,
        })

    extra_rows = 0
    for sku_queue in pending.values():
        while sku_queue:
            item = sku_queue.popleft()
            extra_rows += 1
            anomalies.append({
                "type": "not_in_google_sheets",
                "row": item["row"],
                "sku": item["sku"],
                "quantity": item["quantity"],
                "message": "SKU presente nel file ma non nella giacenza Google Sheets.",
            })

    imported_at = (
        batch.imported_at.isoformat()
        if batch.imported_at is not None
        else None
    )
    return {
        "mapping_required": False,
        "mode": "stock_verification",
        "source": {
            "filename": batch.filename,
            "sheet_name": batch.sheet_name or "",
            "imported_at": imported_at,
        },
        "file": {
            "sheet_name": parsed_file["sheet_name"],
            "header_row": parsed_file["header_row"],
            "mapping": parsed_file["mapping"],
            "parsed_rows": parsed_file["parsed_rows"],
        },
        "summary": {
            "reference_rows": reference_rows,
            "file_rows": parsed_file["parsed_rows"],
            "matched_rows": matched_rows,
            "retained_rows": retained_rows,
            "extra_rows": extra_rows,
            "duplicate_skus": parsed_file.get("duplicate_skus", 0),
            "anomaly_count": len(anomalies),
        },
        "rows": rows,
        "anomalies": anomalies,
    }


def analyze_stock_verification(
    db,
    content: bytes,
    *,
    sku_column: str | None = None,
    quantity_column: str | None = None,
    header_row: int | None = None,
) -> dict:
    parsed = parse_stock_verification_file(
        content,
        sku_column=sku_column,
        quantity_column=quantity_column,
        header_row=header_row,
    )
    if parsed["mapping_required"]:
        return parsed
    return build_stock_verification(db, parsed)


def export_stock_verification(db, verification: dict) -> bytes:
    mappings = load_column_mappings(db)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _excel_sheet_title(
        verification["source"].get("sheet_name")
    )
    headers = [
        mappings["mapping_sku"],
        mappings["mapping_desc"],
        mappings["mapping_lotto"],
        mappings["mapping_qty"],
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    fallback_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for item in verification["rows"]:
        if item["is_spacer"]:
            sheet.append([None, None, None, None])
            continue
        sheet.append([
            item["sku"],
            item["description"],
            item["lotto"],
            item["quantity"],
        ])
        quantity_cell = sheet.cell(sheet.max_row, 4)
        quantity_cell.number_format = "0"
        if item["source"] == "google_sheets":
            quantity_cell.fill = fallback_fill

    sheet.freeze_panes = "A2"
    for index, width in enumerate((22, 54, 28, 14), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    anomaly_sheet = workbook.create_sheet("Anomalie")
    anomaly_headers = ["Tipo", "SKU", "Riga file", "Quantità", "Dettaglio"]
    anomaly_sheet.append(anomaly_headers)
    for cell in anomaly_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for anomaly in verification["anomalies"]:
        anomaly_sheet.append([
            anomaly.get("type", ""),
            anomaly.get("sku", ""),
            anomaly.get("row"),
            anomaly.get("quantity"),
            anomaly.get("message", ""),
        ])
    anomaly_sheet.freeze_panes = "A2"
    for index, width in enumerate((24, 24, 14, 14, 72), start=1):
        anomaly_sheet.column_dimensions[get_column_letter(index)].width = width

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def stock_verification_filename(original_filename: str) -> str:
    safe_stem = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        str(original_filename or "verifica").rsplit(".", 1)[0],
    ).strip("_") or "verifica"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"verifica_giacenze_{safe_stem}_{timestamp}.xlsx"
